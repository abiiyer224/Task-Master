from flask import Flask, render_template, url_for, request, redirect, flash
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime
from flask_migrate import Migrate
import pytz
import pandas as pd
from flask import send_file
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment

#from app import app, db


app=Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///last.db'
db=SQLAlchemy(app)
app.secret_key = "supersecretkey"
migrate = Migrate(app, db)



class Todo(db.Model):
    id=db.Column(db.Integer,primary_key=True)
    content=db.Column(db.String(200),nullable=False)
    date_created = db.Column(db.DateTime, default=lambda: datetime.now(pytz.timezone('Asia/Dubai')))
    completed = db.Column(db.Boolean, default=False)
    completed_date = db.Column(db.DateTime, nullable=True)  # Stores completion time
    #date_completed=db.Column(db.DateTime, default= datetime.utcnow)
    def __repr(self):
        return '<Task %r>' % self.id
    
@app.route('/', methods=['POST','GET'])
def index():
    if request.method =='POST':
        task_content=request.form['content']
        new_task=Todo(content=task_content)
        try:
            db.session.add(new_task)
            db.session.commit()
            return redirect('/')
        except:
            return 'There was an issue adding your task'
    else:
        tasks=Todo.query.order_by(Todo.date_created).all()
        return render_template('index.html',tasks=tasks)
@app.route('/delete/<int:id>')
def delete(id):
    task_to_delete=Todo.query.get_or_404(id)
    try:
        db.session.delete(task_to_delete)
        db.session.commit()
        return redirect('/')
    except:
        return 'There was an issue deleting your task'
    
@app.route('/update/<int:id>',methods=['GET','POST'])
def update(id):
    task= Todo.query.get_or_404(id)
    if request.method =='POST':
        task.content=request.form['content']
        try:
            db.session.commit()
            return redirect('/')
        except:
            return 'There was an issue updating your task'
    else:
        return render_template('update.html', task=task)

@app.route('/complete/<int:id>')
def complete_task(id):
    task= Todo.query.get_or_404(id)
    try:
        if task:
            task.completed = True
            task.completed_date = datetime.utcnow()
            local_timezone = pytz.timezone('Asia/Dubai')
            task.completed_date = datetime.now(local_timezone)
            db.session.commit()
            
        else:
            return 'There was an issue updating your task'
    except:
        return 'There was an issue updating your task'
    return redirect('/')

@app.route('/export')
def export_to_excel():
    tasks = Todo.query.all()
    task_data = []
    for task in tasks:
        task_data.append({
            'ID':task.id,
            'Content':task.content,
            'Date Created':task.date_created.date(),
            'Time Created':task.date_created.strftime('%H:%M:%S'),
            'Completed':task.completed,
            'Completed Date':task.completed_date.date() if task.completed_date else None,
            'Completed Time':task.completed_date.strftime('%H:%M:%S') if task.completed_date else None
        })
    df = pd.DataFrame(task_data)
    excel_file = 'taskmaster.xlsx'
    df.to_excel(excel_file, index=False, engine='openpyxl')
    workbook = load_workbook(excel_file)
    worksheet = workbook.active
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    centre = Alignment(horizontal='center', vertical='center')
    for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
        for cell in row:
            cell.border = thin_border
            cell.alignment = centre
    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter  # Get the column letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)  # Add some padding
        worksheet.column_dimensions[column_letter].width = adjusted_width
    workbook.save(excel_file)
    return send_file(excel_file, as_attachment=True)
if __name__ == "__main__":
    app.run(debug=True)